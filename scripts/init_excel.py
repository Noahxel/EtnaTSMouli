#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to initialize Excel file with students and groupIds from cloned repos
This extracts groupIds from GitLab project names and creates the base Excel structure
"""

import os
import sys
import re
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def extract_student_info_from_name(project_name):
    """
    Extract student name and groupId from GitLab project name
    Format: "Groupe de nom_eleve 1065535" -> ("nom_eleve", 1065535)
    
    Args:
        project_name (str): GitLab project name
        
    Returns:
        tuple: (student_name, group_id) or (None, None) if no match
    """
    # Pattern: "Groupe de student_name digits"
    match = re.search(r'Groupe de ([a-z_]+)\s+(\d+)', project_name, re.IGNORECASE)
    if match:
        return match.group(1), int(match.group(2))
    
    return None, None

def scan_cloned_repos(day):
    """
    Scan cloned repos directory to find student folders
    
    Args:
        day (str): Day number (e.g., 'day01')
        
    Returns:
        list: List of student folder names
    """
    repos_dir = f"./test-repos/{day}/repos"
    if not os.path.exists(repos_dir):
        print(f"❌ Directory not found: {repos_dir}")
        return []
    
    students = []
    for item in os.listdir(repos_dir):
        item_path = os.path.join(repos_dir, item)
        if os.path.isdir(item_path):
            students.append(item)
    
    return sorted(students)

def read_clone_log(day):
    """
    Read clone log to extract groupIds from project names
    The log should contain lines like: "[1/80] Groupe de student_name 1065535"
    
    Args:
        day (str): Day number (e.g., 'day01')
        
    Returns:
        dict: {student_name: group_id}
    """
    log_file = f"./test-repos/{day}/clone_log.txt"
    student_groups = {}
    
    if os.path.exists(log_file):
        with open(log_file, 'r') as f:
            for line in f:
                student_name, group_id = extract_student_info_from_name(line)
                if student_name and group_id:
                    student_groups[student_name] = group_id
    
    return student_groups

def read_config(day):
    """
    Read exercise configuration to get exercise list
    
    Args:
        day (str): Day number (e.g., 'day01')
        
    Returns:
        dict: Configuration data
    """
    config_file = f"./test-repos/{day}/config.json"
    if not os.path.exists(config_file):
        print(f"❌ Config file not found: {config_file}")
        return None
    
    with open(config_file, 'r') as f:
        return json.load(f)

def create_excel_structure(day, students, student_groups, config):
    """
    Create or update Excel file with students, exercises, and groupIds
    Preserves existing test results if Excel already exists
    
    Args:
        day (str): Day number (e.g., 'day01')
        students (list): List of student names
        student_groups (dict): Mapping of student name to groupId
        config (dict): Exercise configuration
    """
    excel_file = './results.xlsx'
    
    # Load existing workbook or create new one
    if os.path.exists(excel_file):
        try:
            wb = Workbook()
            from openpyxl import load_workbook
            wb = load_workbook(excel_file)
            print(f"✅ Loading existing Excel file")
        except:
            wb = Workbook()
            print(f"✅ Creating new Excel file")
    else:
        wb = Workbook()
        if wb.active.title == 'Sheet':
            wb.remove(wb.active)
        print(f"✅ Creating new Excel file")
    
    # Get or create worksheet
    if day in wb.sheetnames:
        ws = wb[day]
        print(f"✅ Updating existing worksheet: {day}")
        # Read existing data - DO NOT DELETE ROWS
        existing_students = set()
        groupid_col_idx = 0
        
        # Find GroupID column
        header_row = ws[1]
        for idx, cell in enumerate(header_row, 1):
            if cell.value == 'GroupID':
                groupid_col_idx = idx
                break
        
        # Get list of existing students
        for row in ws.iter_rows(min_row=2):
            student_id = row[0].value
            if student_id:
                existing_students.add(student_id)
        
        # We'll UPDATE existing students and ADD new ones
        # DO NOT delete any rows - preserve all test results
    else:
        ws = wb.create_sheet(day)
        existing_students = set()
        groupid_col_idx = 0
        print(f"✅ Creating new worksheet: {day}")
    
    # Build headers: Student ID | Score | Ex00 | Ex01 | ... | Ex20 | GroupID
    headers = ['Student ID', 'Score']
    
    # Add setup exercise
    if 'setupExercise' in config:
        headers.append(f"Ex{config['setupExercise']['id']}")
    
    # Add regular exercises
    if 'exercises' in config:
        for ex in config['exercises']:
            headers.append(f"Ex{ex['id']}")
    
    # Add GroupID column at the end
    headers.append('GroupID')
    
    # Only create/update headers if this is a new worksheet
    if not existing_students:
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Set column widths
        ws.column_dimensions['A'].width = 20  # Student ID
        ws.column_dimensions['B'].width = 25  # Score
        for col_idx in range(3, len(headers)):  # Exercises
            col_letter = chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)
            ws.column_dimensions[col_letter].width = 15
        # GroupID column
        groupid_col_letter = chr(64 + len(headers)) if len(headers) <= 26 else 'A' + chr(64 + len(headers) - 26)
        ws.column_dimensions[groupid_col_letter].width = 12
    else:
        # Find GroupID column index for updates
        if groupid_col_idx == 0:
            groupid_col_idx = len(headers)
    
    # Add/update students
    students_added = 0
    students_updated = 0
    groupids_found = len([g for g in student_groups.values() if g])
    
    for student in students:
        group_id = student_groups.get(student, "")
        
        if student in existing_students:
            # FIND and UPDATE existing student row - DO NOT touch test results
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                if row[0].value == student:
                    # Only update GroupID if it's currently empty
                    groupid_cell = ws.cell(row=row_idx, column=groupid_col_idx)
                    if not groupid_cell.value and group_id:
                        groupid_cell.value = group_id
                        groupid_cell.font = Font(color="006100")
                        groupid_cell.alignment = Alignment(horizontal='center', vertical='center')
                    students_updated += 1
                    break
        else:
            # Add NEW student row
            new_row_num = ws.max_row + 1
            
            # Student ID
            ws.cell(row=new_row_num, column=1).value = student
            ws.cell(row=new_row_num, column=1).alignment = Alignment(horizontal='left', vertical='center')
            
            # Score (empty for now)
            ws.cell(row=new_row_num, column=2).value = None
            ws.cell(row=new_row_num, column=2).alignment = Alignment(horizontal='center', vertical='center')
            
            # Exercise columns (empty for now) - NO RED BACKGROUND
            for col in range(3, len(headers)):
                cell = ws.cell(row=new_row_num, column=col)
                cell.value = None  # Explicitly None, not empty string
                cell.alignment = Alignment(horizontal='center', vertical='center')
                # No fill - cells should be white/transparent
            
            # GroupID
            groupid_cell = ws.cell(row=new_row_num, column=len(headers))
            groupid_cell.value = group_id if group_id else None
            groupid_cell.alignment = Alignment(horizontal='center', vertical='center')
            if group_id:
                groupid_cell.font = Font(color="006100")
            
            students_added += 1
    
    # Save workbook
    wb.save(excel_file)
    print(f"✅ Excel file saved: {excel_file}")
    print(f"   - Worksheet: {day}")
    print(f"   - Students added: {students_added}")
    print(f"   - Students updated: {students_updated}")
    print(f"   - Total students: {len(students)}")
    print(f"   - Exercises: {len(headers) - 3}  (excluding Student ID, Score, GroupID)")
    print(f"   - GroupIDs found: {groupids_found}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python init_excel.py <day_number>")
        print("Example: python init_excel.py 1")
        sys.exit(1)
    
    day_num = sys.argv[1]
    day = f"day{int(day_num):02d}"
    
    print(f"🔍 Initializing Excel for {day}...")
    print()
    
    # Read configuration
    config = read_config(day)
    if not config:
        sys.exit(1)
    
    # Scan for students
    students = scan_cloned_repos(day)
    if not students:
        print(f"❌ No student repos found in test-repos/{day}/repos/")
        sys.exit(1)
    
    print(f"✅ Found {len(students)} students")
    
    # Try to read groupIds from log
    student_groups = read_clone_log(day)
    if student_groups:
        print(f"✅ Found {len(student_groups)} groupIds from clone log")
    else:
        print("⚠️  No clone log found - GroupID column will be empty")
        print("   You can manually add groupIds or re-run cloning with logging")
    
    # Create Excel
    create_excel_structure(day, students, student_groups, config)
    print()
    print("📋 Next steps:")
    print(f"   1. Run tests: ./docker-test.sh -d {day} --all")
    print("   2. Check results in results.xlsx")
    print("   3. Validate on ETNA: npm run validate-etna")

if __name__ == '__main__':
    main()
